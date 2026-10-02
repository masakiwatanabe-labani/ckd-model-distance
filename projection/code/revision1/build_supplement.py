# -*- coding: utf-8 -*-
"""Supplementary_Notes.docx（図の legend を含む）と Supplementary_Tables.xlsx を作る。"""
from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

HERE = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(HERE / "code" / "revision1"))
import docx_lib as D  # noqa: E402
import ref_keys as RK  # noqa: E402

SRC = Path.home() / "Downloads" / "C_MS_clean.docx"
FIGDIR = HERE / "manuscript_C" / "figures"
REL_IMAGE = ("http://schemas.openxmlformats.org/officeDocument/2006/"
             "relationships/image")
# 版面幅。pgSz 12240 − pgMar left 1080 − right 1080 = 10080 twips = 7.00 in
TEXT_WIDTH_IN = 7.00


SUB = "@@SUB@@"                      # 小見出しの行に付ける印
SUBHEAD = re.compile(r"^Block \d+\.\s")


def figure_file(n: int) -> Path:
    """Figure S<n> の画像。番号は本文の初出順で振ってあるので接頭辞で引く。"""
    cand = sorted(FIGDIR.glob(f"FigS{n}_*.png"))
    assert len(cand) == 1, f"FigS{n} の画像が {len(cand)} 件: {[c.name for c in cand]}"
    return cand[0]


def png_size(p: Path) -> tuple[tuple[int, int], float]:
    """PNG の画素数と dpi。dpi が無ければ 300 とみなす。"""
    from PIL import Image
    im = Image.open(p)
    dpi = im.info.get("dpi", (300.0, 300.0))[0] or 300.0
    return im.size, float(dpi)

DEST = Path.home() / "Desktop" / "CKD_xspecies_submission"
SUP = HERE / "manuscript_C" / "SUPPLEMENTARY.md"


def parse():
    """補足 Markdown を Notes / Tables / Figure legends に分ける。"""
    notes, tables, figs, cur = [], {}, {}, None
    for blk in SUP.read_text().split("\n\n"):
        s = blk.strip()
        if not s:
            continue
        if s.startswith("# ") or s.startswith("## "):
            cur = None
            continue
        mt, mf = re.match(r"^Table S(\d+)\.", s), re.match(r"^Figure S(\d+)\.", s)
        if mt:
            cur = ("T", int(mt.group(1)))
            tables[cur[1]] = {"caption": s, "rows": []}
            continue
        if mf:
            cur = ("F", int(mf.group(1)))
            figs[cur[1]] = s
            continue
        if cur and cur[0] == "T":
            if s.lstrip().startswith("|"):
                tables[cur[1]]["rows"] += s.splitlines()
            elif SUBHEAD.match(s):
                # 「Block 1. …」のような小見出しは、続く表の見出し行の直前に置く行。
                # 説明文に足してはいけない（BL2 で見つかった流れ込みの原因）。
                tables[cur[1]]["rows"].append(SUB + s)
            elif tables[cur[1]]["rows"]:
                raise SystemExit(
                    f"Table S{cur[1]}: 表が始まったあとの地の文は説明文に足せない: {s[:70]}")
            else:
                tables[cur[1]]["caption"] += " " + s
        elif cur and cur[0] == "F":
            figs[cur[1]] += " " + s
        else:
            notes.append(s)
    return notes, tables, figs


def cells(rows):
    """表の行を升目に開く。小見出しの行は 1 升だけの行として返す。"""
    out = []
    for r in rows:
        if r.startswith(SUB):
            out.append([r[len(SUB):]])
            continue
        if re.fullmatch(r"\|[\s\-|:]+\|", r.strip()):
            continue
        out.append([c.strip().replace("\\|", "|")
                    for c in re.split(r"(?<!\\)\|", r.strip().strip("|"))])
    return out


def _li_ref() -> str:          # noqa: D401  （ref_keys に委譲）
    return RK.table()["li"]


def _li_ref_unused() -> str:
    """Li ら（ネコ CKD）の文献番号を本文の一覧から引く。番号を振り直しても追随する。"""
    import re as _re
    for line in (HERE / "manuscript_C" / "REFERENCES.md").read_text().splitlines():
        m = _re.match(r"^(\d+)\.\s+(.*)$", line.strip())
        if m and "Integrated multi-omics analysis of renal metabolism in domestic cats" in m.group(2):
            return m.group(1)
    raise SystemExit("Li et al. の文献が見つからない")


def main() -> int:
    notes, tables, figs = parse()
    # [REF:key] を現在の番号に解く（番号を書き込んだ箇所を残さない）
    notes = [RK.resolve(s) for s in notes]
    for n in tables:
        tables[n]["caption"] = RK.resolve(tables[n]["caption"])
        tables[n]["rows"] = [RK.resolve(r) for r in tables[n]["rows"]]
    figs = {n: RK.resolve(v) for n, v in figs.items()}
    DEST.mkdir(parents=True, exist_ok=True)

    # ---------------------------------------------------------- Notes docx
    body = [D.para("Supplementary Materials", bold=True)]
    body.append(D.para("Mean-Centering Alters Pathway-Based Transcriptomic Comparisons of "
                       "Feline and Mouse Kidney Disease"))
    body.append(D.para("Masaki Watanabe, Takeru Sasaki, Ryuya Nakagawa and Nobuya Sasaki"))
    body.append(D.para(""))
    for s in notes:
        body.append(D.para(s, bold=s.startswith("Supplementary Note S")))
    body.append(D.para(""))
    body.append(D.para("Supplementary Tables", bold=True))
    body.append(D.para("The tables themselves are supplied as Supplementary_Tables.xlsx, one "
                       "sheet per table. Their captions are reproduced here."))
    for n in sorted(tables):
        body.append(D.para(tables[n]["caption"]))
    body.append(D.para(""))
    body.append(D.para("Supplementary Figures", bold=True))
    body.append(D.para("Each figure is placed above its legend. The images are the same files "
                       "as the uploaded figure originals."))
    media, fig_rels = {}, []
    for n in sorted(figs):
        img = figure_file(n)
        rid = f"rIdFigS{n}"
        part = f"word/media/figS{n}{img.suffix}"
        media[part] = img.read_bytes()
        fig_rels.append(f'<Relationship Id="{rid}" Type="{REL_IMAGE}" '
                        f'Target="media/figS{n}{img.suffix}"/>')
        px, dpi = png_size(img)
        body.append(D.image_para(rid, px, dpi, TEXT_WIDTH_IN, doc_id=900 + n,
                                 name=f"Figure S{n}",
                                 descr=figs[n].split(". ")[1] if ". " in figs[n] else f"Figure S{n}"))
        body.append(D.para(figs[n]))
        print(f"  Figure S{n}: {img.name} {px[0]}x{px[1]} px @ {dpi:.0f} dpi -> "
              f"{min(px[0] / dpi, TEXT_WIDTH_IN):.2f} in 幅")

    xml = zipfile.ZipFile(SRC).read("word/document.xml").decode("utf-8")
    sect = re.search(r"<w:sectPr[ >].*?</w:sectPr>", xml, re.S)
    head = xml[:xml.index("<w:body>") + len("<w:body>")]
    new = head + "".join(body) + (sect.group(0) if sect else "") + "</w:body></w:document>"
    core = zipfile.ZipFile(SRC).read("docProps/core.xml").decode("utf-8")
    core = re.sub(r"<dc:title>.*?</dc:title>", "<dc:title>Supplementary Notes</dc:title>", core)
    for tag in ("dc:creator", "cp:lastModifiedBy", "dc:subject", "dc:description"):
        core = re.sub(rf"<{tag}>.*?</{tag}>", f"<{tag}></{tag}>", core)
    _z = zipfile.ZipFile(SRC)
    _media = tuple(n for n in _z.namelist() if n.startswith("word/media/"))
    _rels = re.sub(r'<Relationship[^>]*Target="media/[^"]*"[^>]*/>', "",
                   _z.read("word/_rels/document.xml.rels").decode("utf-8"))
    _rels = _rels.replace("</Relationships>", "".join(fig_rels) + "</Relationships>")
    D.write_docx(SRC, DEST / "Supplementary_Notes.docx", new,
                 drop=("word/comments.xml", "word/commentsExtended.xml",
                       "word/commentsIds.xml") + _media,
                 extra={"docProps/core.xml": core,
                        "word/_rels/document.xml.rels": _rels},
                 add=media)
    print(f"  Supplementary_Notes.docx: Notes {len(notes)} ブロック / "
          f"表キャプション {len(tables)} / 図 legend {len(figs)}")

    # -------------------------------------------- Figures PDF（図の原図、現在の番号順）
    # 投稿済みの補足図 PDF は旧番号の並び（S1 = pathway alignment）
    # なので、本文の番号と合うように作り直す。
    from pypdf import PdfWriter
    wr = PdfWriter()
    for n in sorted(figs):
        src = figure_file(n).with_suffix(".pdf")
        assert src.exists(), f"{src} が無い"
        wr.append(str(src))
    wr.add_metadata({"/Title": "Supplementary Figures S1-S5", "/Author": "",
                     "/Creator": "", "/Producer": "", "/Subject": "", "/Keywords": ""})
    with open(DEST / "Supplementary_Figures.pdf", "wb") as fh:
        wr.write(fh)
    print(f"  Supplementary_Figures.pdf: {len(figs)} ページ（Figure S1-S{max(figs)} の順）")

    # ---------------------------------------------------------- Tables xlsx
    wb = Workbook()
    wb.remove(wb.active)
    for n in sorted(tables):
        ws = wb.create_sheet(f"Table S{n}")
        ws["A1"] = tables[n]["caption"]
        ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
        ws["A1"].font = Font(size=9)
        ws.row_dimensions[1].height = 58
        data = cells(tables[n]["rows"])
        # 太字にするのは、小見出しの行と、各ブロックの列見出し行（小見出しの次の行、
        # および表の先頭行）。
        head_at = {0} if data and len(data[0]) > 1 else set()
        subs = set()
        for i, row in enumerate(data):
            if len(row) == 1:
                subs.add(i)
                if i + 1 < len(data):
                    head_at.add(i + 1)
        for r, row in enumerate(data, start=3):
            for c, v in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c, value=_num(v))
                if (r - 3) in subs or (r - 3) in head_at:
                    cell.font = Font(bold=True)
        w = max((len(str(x)) for row in data for x in row), default=12)
        for c in range(1, (max((len(r) for r in data), default=1)) + 1):
            ws.column_dimensions[chr(64 + c)].width = min(max(12, w), 40)
        ws.freeze_panes = "A4"
    # Table S10 の per-gene 検出コールは行数が多いので、別シートで丸ごと付ける（AF3）
    calls = HERE / "results" / "roundR" / "groupA_detection_calls.tsv"
    if calls.exists():
        import csv
        # 検出コールの表番号は振り直しで動くので、キャプションから読む
        mcap = re.search(r"(?m)^Table S(\d+)\. Proteomic detection calls",
                         (HERE / "manuscript_C" / "SUPPLEMENTARY.md").read_text())
        tno = mcap.group(1) if mcap else "?"
        ws = wb.create_sheet(f"Table S{tno} gene list")
        ws["A1"] = (f"Per-gene proteomic detection calls behind the Group A selection "
                    f"(Table S{tno}). Detection requires a quantified value in at least half of "
                    "the samples of the dataset concerned: at least 12 of 23 feline cortical and "
                    f"10 of 19 feline medullary samples from the published supplementary data of "
                    f"Li et al. [{_li_ref()}] (42003_2025_9164_MOESM3_ESM.xlsx; raw data in "
                    "ProteomeXchange under PXD066590), where an unquantified sample is blank, "
                    "and at least 5 of 9 Pod-TRECK samples, where an unquantified sample is "
                    "recorded as 0 and a zero is therefore read as not quantified. Rows are the "
                    "genes of the Δ matrix.")
        ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
        ws["A1"].font = Font(size=9)
        ws.row_dimensions[1].height = 42
        with open(calls) as fh:
            for r, row in enumerate(csv.reader(fh, delimiter="\t"), start=3):
                for c, v in enumerate(row, start=1):
                    cell = ws.cell(row=r, column=c, value=v)
                    if r == 3:
                        cell.font = Font(bold=True)
        for c, w in zip("ABCDEFG", (14, 22, 23, 20, 8, 22, 22)):
            ws.column_dimensions[c].width = w
        ws.freeze_panes = "A4"
        print(f"  Table S{tno} gene list: {ws.max_row - 3} 行")

    wb.save(DEST / "Supplementary_Tables.xlsx")
    tot = sum(len(cells(t["rows"])) for t in tables.values())
    print(f"  Supplementary_Tables.xlsx: {len(tables)} シート / データ行 {tot}")
    return 0


def _num(v: str):
    s = str(v).replace("−", "-").replace(",", "")
    try:
        return int(s) if re.fullmatch(r"-?\d+", s) else float(s)
    except ValueError:
        return v


if __name__ == "__main__":
    sys.exit(main())
