# -*- coding: utf-8 -*-
"""BL: 出来上がった補足ファイルの説明文を点検する。

  1. どの説明文にも、同じ文の繰り返しが無いこと
  2. どの説明文にも、表の中に置くべき小見出し（「Block n. …」）が入り込んでいないこと
  3. 補足ノートと補足表の説明文が、表ごとに一字一句一致すること
  4. 同じ小見出しが表の中で重複していないこと
  5. 本文が述べる図の提出形態が、実際の提出物と合っていること（BL1）
"""
from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parents[2]
M = HERE / "manuscript_C"
DEST = Path.home() / "Desktop" / "CKD_xspecies_submission"
NOTES = DEST / "Supplementary_Notes.docx"
XLSX = DEST / "Supplementary_Tables.xlsx"
FIGPDF = DEST / "Supplementary_Figures.pdf"
SUBHEAD = re.compile(r"Block \d+\.\s+[A-Z]")

FAIL: list[str] = []
OK = 0


def chk(what, ok, detail=""):
    global OK
    if ok:
        OK += 1
    else:
        FAIL.append(f"{what}{(' — ' + detail) if detail else ''}")
    print(f"  {'ok ' if ok else '✗  '} {what}{(' — ' + detail) if detail else ''}")


def notes_captions() -> dict[str, str]:
    x = zipfile.ZipFile(NOTES).read("word/document.xml").decode()
    out = {}
    for b in re.findall(r"<w:p\b.*?</w:p>", x, re.S):
        t = " ".join(re.sub(r"<[^>]+>", "", b).split())
        m = re.match(r"^(Table S\d+|Figure S\d+)\.", t)
        if m:
            out[m.group(1)] = t
    return out


def xlsx_sheets():
    from openpyxl import load_workbook
    wb = load_workbook(XLSX, read_only=True)
    cap, rows = {}, {}
    for ws in wb.worksheets:
        data = list(ws.iter_rows(values_only=True))
        if not data:
            continue
        name = ws.title
        if isinstance(data[0][0], str):
            cap[name] = " ".join(str(data[0][0]).split())
        rows[name] = [[c for c in r if c is not None] for r in data[1:]]
    return cap, rows


def sentences(s: str) -> list[str]:
    return [x.strip() for x in re.split(r"(?<=[.])\s+", s) if x.strip()]


def main() -> int:
    for p in (NOTES, XLSX):
        if not p.exists():
            print(f"無い: {p}", file=sys.stderr)
            return 2
    ncap = notes_captions()
    xcap, xrows = xlsx_sheets()

    # 1 と 2
    dup, leaked = [], []
    for k, v in sorted(ncap.items()):
        ss = sentences(v)
        if len(ss) != len(set(ss)):
            dup.append(f"{k}: " + [x for x in ss if ss.count(x) > 1][0][:60])
        if SUBHEAD.search(v):
            leaked.append(f"{k}: " + SUBHEAD.search(v).group(0))
    chk(f"説明文 {len(ncap)} 件に同じ文の繰り返しが無い", not dup, "; ".join(dup[:3]))
    chk("説明文に表の小見出し（Block n. …）が入り込んでいない", not leaked,
        "; ".join(leaked[:3]))

    # 3
    diff = []
    for k, v in sorted(xcap.items()):
        if k.endswith("gene list"):
            continue
        if k not in ncap:
            diff.append(f"{k}: 補足ノートに説明文が無い")
        elif ncap[k] != v:
            diff.append(f"{k}: ノートと表で違う")
    chk(f"補足ノートと補足表の説明文 {len(xcap) - 1} 件が一致", not diff, "; ".join(diff[:3]))

    # 4
    bad_sub = []
    for name, rows in sorted(xrows.items()):
        subs = [r[0] for r in rows if len(r) == 1 and SUBHEAD.match(str(r[0]))]
        if len(subs) != len(set(subs)):
            bad_sub.append(f"{name}: {subs}")
        for s in subs:
            n = re.match(r"Block (\d+)\.", s).group(1)
            if sum(1 for t in subs if t.startswith(f"Block {n}.")) > 1:
                bad_sub.append(f"{name}: Block {n}. が複数")
    chk("表の中の小見出しに重複が無い", not bad_sub, "; ".join(sorted(set(bad_sub))[:3]))
    n_sub = sum(1 for rows in xrows.values() for r in rows
                if len(r) == 1 and SUBHEAD.match(str(r[0])))
    print(f"      （表の中の小見出し {n_sub} 件）")

    # 5: BL1 本文の提出形態
    back = " ".join((M / "BACKMATTER.md").read_text().split())
    say_pdf = ("single PDF file (Supplementary_Figures.pdf)" in back
               and "shown above their legends in the Supplementary Notes" in back)
    chk("本文が図を「1 つの PDF + 補足ノートに掲載」と述べている", say_pdf)
    chk("その PDF が実際にある", FIGPDF.exists(), FIGPDF.name)
    media = [n for n in zipfile.ZipFile(NOTES).namelist() if n.startswith("word/media/")]
    chk("補足ノートに図が 5 点入っている", len(media) == 5, f"{len(media)} 点")
    chk("「supplied as separate files」が残っていない",
        "supplied as separate files" not in back)
    # 回答書の自主訂正の項目 8 と矛盾しないこと
    lt = (M / "RESPONSE.md")
    if lt.exists():
        t = " ".join(lt.read_text().split())
        chk("回答書の項目 8 も「S1 to S5 の順」「補足ノートに掲載」と述べている",
            "in the order S1 to S5" in t
            and "embedded above their legends in the Supplementary Notes" in t)

    print(f"\n合格 {OK} / 不合格 {len(FAIL)}")
    for f in FAIL:
        print("  ✗", f)
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
