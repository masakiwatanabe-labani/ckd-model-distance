# -*- coding: utf-8 -*-
"""BJ2: 提出する補足・回答書に現れる文献番号が、いまの文献一覧と合っているか。

本文の番号は renumber_refs_C.py が振り直すが、補足ノート・補足表・回答書・カバー
レターは別に作られるので、文献が増減すると古い番号が残りうる。ここでは
**出来上がったファイルの文字**から「[n]」「Reference n」を拾い、

  - その番号が文献一覧にあること
  - 文中の手がかり（Li ら、GSE98622、ssGSEA …）と、引かれた文献の書誌が合うこと
  - 鍵（ref_keys）で引くべき箇所が、鍵から引いた番号になっていること

を確かめる。査読票の引用に出る番号は投稿時の番号なので対象から外す。
"""
from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(HERE / "code" / "revision1"))
import ref_keys as RK  # noqa: E402
from ay_check_citations import CLUES  # noqa: E402

DEST = Path.home() / "Desktop" / "CKD_xspecies_submission"
M = HERE / "manuscript_C"
FILES = ["Supplementary_Notes.docx", "Supplementary_Tables.xlsx",
         "Response_to_reviewers.docx", "Cover_letter.docx"]
CITE = re.compile(r"\[(\d+(?:\s*[,–-]\s*\d+)*)\]|\b[Rr]eference (\d+)\b")
NOTE_CUE = "follows the submitted numbering"
# 鍵ごとの手がかり。これが当たる文の番号は、鍵から引いた番号でなければならない。
KEY_CUES = {
    "li": r"Li et al\.|42003_2025_9164|domestic cats|feline cortical and medullary",
    "companion": r"companion study|companion paper",
    "iri": r"\bGSE98622\b",
    "ensembl": r"\bEnsembl\b|BioMart",
    "podtreck": r"\bGSE299326\b|Pod-TRECK model",
    "ercb": r"\bGSE104954\b",
}


def docx_text(p: Path) -> str:
    x = zipfile.ZipFile(p).read("word/document.xml").decode()
    return " ".join(re.sub(r"<[^>]+>", " ", re.sub(r"</w:p>", "\n", x)).split())


def xlsx_text(p: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str) and v.strip():
                    out.append(v)
    return "\n".join(out)


def refs() -> dict[str, str]:
    out = {}
    for line in (M / "REFERENCES.md").read_text().splitlines():
        m = re.match(r"^(\d+)\.\s+(.*)$", line.strip())
        if m:
            out[m.group(1)] = m.group(2)
    return out


def _quotes() -> dict:
    """査読コメントの原文（キー → 本文）。"""
    sys.path.insert(0, str(HERE / "code" / "revision1"))
    import build_response_letter as BRL
    return BRL.quotes()


def quoted() -> str:
    """査読票の原文。ここに出る番号は投稿時の番号なので触らない。"""
    p = M / "REVIEWER_COMMENTS.md"
    return " ".join(p.read_text().split()) if p.exists() else ""


def sentence_of(text: str, pos: int) -> str:
    a = text.rfind(". ", 0, pos)
    b = text.find(". ", pos)
    return text[(a + 2 if a >= 0 else 0):(b + 1 if b > 0 else len(text))]


def main() -> int:
    R, Q, T = refs(), quoted(), RK.table()
    bad, skipped, checked = [], 0, 0
    print(f"文献一覧 {len(R)} 件 / 鍵から引いた番号 {T}")
    for name in FILES:
        p = DEST / name
        if not p.exists():
            print(f"  （無い）{name}")
            continue
        text = xlsx_text(p) if p.suffix == ".xlsx" else docx_text(p)
        found = 0
        for m in CITE.finditer(text):
            nums = ([n.strip() for n in re.split(r"[,–-]", m.group(1))]
                    if m.group(1) else [m.group(2)])
            sent = sentence_of(text, m.start())
            flat = " ".join(sent.split())
            if flat and flat[:60] in Q:
                skipped += 1
                continue
            if NOTE_CUE in flat:            # BK の注記は下で別に検査する
                skipped += 1
                continue
            found += 1
            for n in nums:
                checked += 1
                if n not in R:
                    bad.append(f"{name}: [{n}] は文献一覧に無い（1–{max(map(int, R))}）: "
                               f"{flat[:70]}")
                    continue
                # 手がかりが当たる文なら、引かれた文献と合っているか
                for pat, need in CLUES:
                    if re.search(pat, flat) and not any(w.lower() in R[n].lower()
                                                        for w in need):
                        # 同じ文の別の番号が当たっていれば良し
                        if any(any(w.lower() in R[o].lower() for w in need)
                               for o in nums if o in R):
                            continue
                        bad.append(f"{name}: {need[0]} を指すべき文の引用が [{n}]: "
                                   f"{flat[:70]}")
                # 鍵で引くべき箇所の照合（番号側からと、文の手がかり側からの両方）
                for key, needle in RK.NEEDLE.items():
                    if needle[:40] in R[n] and n != T[key]:
                        bad.append(f"{name}: 鍵 {key} は [{T[key]}] のはずが [{n}]")
                for key, cue in KEY_CUES.items():
                    if re.search(cue, flat) and n != T[key] and len(nums) == 1:
                        bad.append(f"{name}: 「{key}」を指す文の番号が [{n}]。"
                                   f"鍵から引くと [{T[key]}]: {flat[:70]}")
        # 回答書・カバーレターは鍵からしか番号を書かない。鍵に無い番号は古い番号。
        if name in ("Response_to_reviewers.docx", "Cover_letter.docx"):
            for m in re.finditer(r"\b[Rr]eference (\d+)\b", text):
                flat = " ".join(sentence_of(text, m.start()).split())
                if flat and flat[:60] in Q:
                    continue
                if m.group(1) not in set(T.values()):
                    bad.append(f"{name}: 「Reference {m.group(1)}」は鍵から引いた番号"
                               f"（{sorted(set(T.values()), key=int)}）に無い: {flat[:64]}")
        print(f"  {name}: 検査対象の番号 {found} 箇所")
    # BK: 査読票の原文に出る文献番号は投稿時の番号。引用は変えず、注記を添える。
    letter = DEST / "Response_to_reviewers.docx"
    if letter.exists():
        txt = docx_text(letter)
        notes = 0
        for key, q in sorted(_quotes().items()):
            if " ".join(q.split())[:60] not in txt:
                continue                     # この回答書に載っていない査読コメント
            for old_n, new_n, _e in RK.quote_renumbering(q):
                notes += 1
                if f"[{old_n}]" not in txt:
                    bad.append(f"回答書: 査読コメント {key} の原文の [{old_n}] が変わっている")
                want = (f"reference [{old_n}] in the submitted version is "
                        f"reference [{new_n}] here")
                if want not in txt:
                    bad.append(f"回答書: {key} の番号の注記が無い、または古い。要: {want}")
        print(f"  査読票の原文に出る文献番号の注記 {notes} 件")
    print(f"\n突き合わせた番号 {checked} 件 / 査読票の引用として除いた {skipped} 件")
    if bad:
        for b in bad[:10]:
            print("  ✗ " + b)
        return 1
    print("  すべて現在の文献一覧と合っている")
    return 0


if __name__ == "__main__":
    sys.exit(main())
